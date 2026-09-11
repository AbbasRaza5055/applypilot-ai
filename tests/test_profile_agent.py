"""Unit tests for AP-001 profile and CV intelligence."""

from __future__ import annotations

from types import SimpleNamespace

import pytest

from agents import profile_agent
from models import CandidateProfile
from services.document_service import (
    EmptyDocumentError,
    UnsupportedResumeFormatError,
    extract_text_from_resume,
)


def _write_pdf_with_text(file_path, text):
    """Create a minimal, valid PDF with extractable Helvetica text."""
    content_stream = f"BT\n/F1 18 Tf\n72 720 Td\n({text}) Tj\nET\n".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content_stream)).encode("ascii") + b" >>\nstream\n"
        + content_stream
        + b"endstream",
    ]

    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_number, object_content in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{object_number} 0 obj\n".encode("ascii"))
        pdf.extend(object_content)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    file_path.write_bytes(pdf)


def test_extract_text_from_pdf(tmp_path):
    resume_path = tmp_path / "resume.pdf"
    _write_pdf_with_text(resume_path, "Jane Doe - Python Developer")

    extracted_text = extract_text_from_resume(str(resume_path))

    assert "Jane Doe" in extracted_text
    assert "Python Developer" in extracted_text


def test_extract_text_from_docx(tmp_path):
    from docx import Document

    resume_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Data analyst")
    document.save(resume_path)

    assert extract_text_from_resume(str(resume_path)) == "Jane Doe\nData analyst"


def test_extract_text_rejects_unsupported_file_type(tmp_path):
    resume_path = tmp_path / "resume.txt"
    resume_path.write_text("Not a supported resume")

    with pytest.raises(UnsupportedResumeFormatError, match="Supported types"):
        extract_text_from_resume(str(resume_path))


def test_extract_text_rejects_empty_document(tmp_path):
    from docx import Document

    resume_path = tmp_path / "empty.docx"
    Document().save(resume_path)

    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        extract_text_from_resume(str(resume_path))


def test_extract_text_rejects_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError, match="does not exist"):
        extract_text_from_resume(str(tmp_path / "missing.pdf"))


def _mock_gemini_client(monkeypatch, parsed_response):
    captured = {}

    class FakeGenerateContentConfig:
        def __init__(self, **kwargs):
            captured["config"] = kwargs

    class FakeModels:
        def generate_content(self, **kwargs):
            captured["request"] = kwargs
            return SimpleNamespace(parsed=parsed_response)

    client = SimpleNamespace(models=FakeModels())
    types = SimpleNamespace(GenerateContentConfig=FakeGenerateContentConfig)
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    monkeypatch.setattr(
        profile_agent, "_get_gemini_client", lambda api_key: (client, types)
    )
    return captured


def test_analyze_candidate_profile_returns_validated_model(monkeypatch):
    captured = _mock_gemini_client(
        monkeypatch,
        {
            "name": "Jane Doe",
            "email": "jane@example.com",
            "skills": ["Python"],
            "years_of_experience": 2.0,
        },
    )

    profile = profile_agent.analyze_candidate_profile("Jane Doe has two years of Python.")

    assert isinstance(profile, CandidateProfile)
    assert profile.name == "Jane Doe"
    assert profile.skills == ["Python"]
    assert profile.education == []
    assert captured["config"]["response_schema"] is CandidateProfile
    assert captured["config"]["response_mime_type"] == "application/json"


def test_analyze_candidate_profile_rejects_invalid_llm_output(monkeypatch):
    _mock_gemini_client(monkeypatch, {"email": "jane@example.com"})

    with pytest.raises(profile_agent.ProfileValidationError, match="CandidateProfile"):
        profile_agent.analyze_candidate_profile("Resume content")


def test_process_resume_returns_candidate_profile(monkeypatch):
    expected_profile = CandidateProfile(name="Jane Doe", skills=["Python"])
    monkeypatch.setattr(
        profile_agent, "extract_text_from_resume", lambda file_path: "Resume content"
    )
    monkeypatch.setattr(
        profile_agent, "analyze_candidate_profile", lambda text: expected_profile)

    assert profile_agent.process_resume("resume.pdf") is expected_profile
