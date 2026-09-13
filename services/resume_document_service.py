from __future__ import annotations

from pathlib import Path


class ResumeDocumentError(RuntimeError):
    """Base exception for resume document generation failures."""


class EmptyResumeError(ResumeDocumentError):
    """Raised when the optimized resume is empty."""


class ResumeDocumentGenerationError(ResumeDocumentError):
    """Raised when a resume document cannot be generated."""


def _validate_resume_text(optimized_resume: str) -> str:
    if not isinstance(optimized_resume, str):
        raise TypeError("optimized_resume must be a string.")

    normalized = optimized_resume.strip()

    if not normalized:
        raise EmptyResumeError("optimized_resume cannot be empty.")

    return normalized


def _prepare_output_path(output_path: str) -> Path:
    if not isinstance(output_path, str):
        raise TypeError("output_path must be a string.")

    path = Path(output_path)

    if not path.name:
        raise ResumeDocumentGenerationError(
            "output_path must include a file name."
        )

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ResumeDocumentGenerationError(
            f"Unable to create output directory for '{path}'."
        ) from exc

    return path


def generate_resume_docx(
    optimized_resume: str,
    output_path: str,
) -> str:
    """Generate a DOCX resume from optimized resume text."""
    text = _validate_resume_text(optimized_resume)
    path = _prepare_output_path(output_path)

    try:
        from docx import Document
        from docx.shared import Pt

        document = Document()

        for line in text.splitlines():
            line = line.strip()

            if not line:
                document.add_paragraph()
                continue

            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.size = Pt(10.5)

        document.save(str(path))

    except ImportError as exc:
        raise ResumeDocumentGenerationError(
            "DOCX generation requires the 'python-docx' package."
        ) from exc

    except Exception as exc:
        raise ResumeDocumentGenerationError(
            f"Unable to generate DOCX resume '{path}'."
        ) from exc

    return str(path)


def generate_resume_pdf(
    optimized_resume: str,
    output_path: str,
) -> str:
    """Generate a PDF resume from optimized resume text."""
    text = _validate_resume_text(optimized_resume)
    path = _prepare_output_path(output_path)

    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

    except ImportError as exc:
        raise ResumeDocumentGenerationError(
            "PDF generation requires the 'reportlab' package."
        ) from exc

    try:
        page_width, page_height = letter

        pdf = canvas.Canvas(
            str(path),
            pagesize=letter,
        )

        margin = 50
        line_height = 15
        x = margin
        y = page_height - margin
        max_width = page_width - (2 * margin)

        pdf.setFont("Helvetica", 10)

        for raw_line in text.splitlines():
            line = raw_line.strip()

            if not line:
                y -= line_height
                if y <= margin:
                    pdf.showPage()
                    pdf.setFont("Helvetica", 10)
                    y = page_height - margin
                continue

            words = line.split()
            current_line = ""

            for word in words:
                candidate_line = (
                    word
                    if not current_line
                    else f"{current_line} {word}"
                )

                if pdf.stringWidth(
                    candidate_line,
                    "Helvetica",
                    10,
                ) <= max_width:
                    current_line = candidate_line
                    continue

                pdf.drawString(x, y, current_line)
                y -= line_height
                current_line = word

                if y <= margin:
                    pdf.showPage()
                    pdf.setFont("Helvetica", 10)
                    y = page_height - margin

            if current_line:
                pdf.drawString(x, y, current_line)
                y -= line_height

            if y <= margin:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = page_height - margin

        pdf.save()

    except Exception as exc:
        raise ResumeDocumentGenerationError(
            f"Unable to generate PDF resume '{path}'."
        ) from exc

    return str(path)