"""Resume document validation and text extraction utilities."""

from __future__ import annotations

from pathlib import Path


class DocumentServiceError(RuntimeError):
    """Base exception for resume document processing failures."""


class UnsupportedResumeFormatError(DocumentServiceError):
    """Raised when a resume is not a PDF or DOCX file."""


class EmptyDocumentError(DocumentServiceError):
    """Raised when a supported resume has no extractable text."""


class DocumentExtractionError(DocumentServiceError):
    """Raised when a supported resume cannot be read."""


SUPPORTED_RESUME_EXTENSIONS = {".pdf", ".docx"}


def _extract_pdf_text(file_path: Path) -> str:
    """Extract text from a PDF while keeping the PDF dependency optional at import time."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentExtractionError(
            "PDF resume extraction requires the 'pypdf' package. "
            "Install it before processing PDF files."
        ) from exc

    try:
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentExtractionError(
            f"Unable to read PDF resume '{file_path}': {exc}"
        ) from exc


def _extract_docx_text(file_path: Path) -> str:
    """Extract paragraph and table text from a DOCX document."""
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentExtractionError(
            "DOCX resume extraction requires the 'python-docx' package. "
            "Install it before processing DOCX files."
        ) from exc

    try:
        document = Document(str(file_path))
        text_parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                text_parts.extend(cell.text for cell in row.cells)
        return "\n".join(text_parts)
    except Exception as exc:
        raise DocumentExtractionError(
            f"Unable to read DOCX resume '{file_path}': {exc}"
        ) from exc


def extract_text_from_resume(file_path: str) -> str:
    """Return readable text from a PDF or DOCX resume.

    The function validates the extension and source file before extraction. It
    raises a descriptive exception instead of accepting unsupported, empty, or
    unreadable documents.
    """
    if not isinstance(file_path, str):
        raise TypeError("file_path must be a string path to a PDF or DOCX resume.")

    resume_path = Path(file_path)
    extension = resume_path.suffix.lower()
    if extension not in SUPPORTED_RESUME_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_RESUME_EXTENSIONS))
        received = extension or "no extension"
        raise UnsupportedResumeFormatError(
            f"Unsupported resume file type '{received}'. Supported types: {supported}."
        )

    if not resume_path.is_file():
        raise FileNotFoundError(f"Resume file does not exist: {resume_path}")

    extracted_text = (
        _extract_pdf_text(resume_path)
        if extension == ".pdf"
        else _extract_docx_text(resume_path)
    )
    normalized_text = extracted_text.strip()
    if not normalized_text:
        raise EmptyDocumentError(
            f"Resume '{resume_path}' contains no extractable text."
        )

    return normalized_text
